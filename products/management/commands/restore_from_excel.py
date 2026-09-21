import openpyxl
from django.core.management.base import BaseCommand
from products.models import Product, Category, Brand, Festival
from decimal import Decimal
import random
from django.utils import timezone


class Command(BaseCommand):
    help = 'Restore products from Excel file (correct current data)'

    def handle(self, *args, **options):
        self.stdout.write('Restoring products from Excel file...')
        
        try:
            wb = openpyxl.load_workbook('Vamsi_Crackers 2026 diwali.xlsx')
            ws = wb.active
            
            # Clear existing data
            Product.objects.all().delete()
            Category.objects.all().delete()
            self.stdout.write('Cleared existing data')
            
            # Get unique categories from Excel
            categories = set()
            for row in ws.iter_rows(min_row=2, values_only=True):
                if row and len(row) >= 2:
                    category_name = str(row[1]).strip() if row[1] else ''
                    if category_name:
                        categories.add(category_name)
            
            # Create categories
            category_map = {}
            for i, cat_name in enumerate(sorted(categories)):
                slug = cat_name.lower().replace(' ', '-').replace('/', '-').replace('&', 'and').replace('(', '').replace(')', '').replace(',', '')
                category = Category.objects.create(
                    name=cat_name,
                    slug=slug,
                    description=f'{cat_name} products',
                    order=i + 1,
                    is_active=True
                )
                category_map[cat_name] = category
                self.stdout.write(f'Created category: {cat_name}')
            
            # Create brand
            brand, _ = Brand.objects.get_or_create(
                name='RedApple',
                defaults={'slug': 'redapple', 'is_featured': True}
            )
            
            # Create festival
            festival, _ = Festival.objects.get_or_create(
                name='Diwali',
                defaults={
                    'slug': 'diwali',
                    'description': 'Festival of Lights',
                    'start_date': timezone.now().date(),
                    'end_date': timezone.now().date(),
                    'is_active': True
                }
            )
            
            # Create products from Excel
            created_count = 0
            for index, row in enumerate(ws.iter_rows(min_row=2, values_only=True)):
                if not row or len(row) < 3:
                    continue
                
                product_name = str(row[0]).strip() if row[0] else ''
                category_name = str(row[1]).strip() if row[1] else ''
                original_price = row[2] if row[2] else 0
                offer_price = row[3] if len(row) > 3 and row[3] else None
                
                if not product_name or not category_name or original_price == 0:
                    continue
                
                # Apply 10% markup
                base_price = offer_price if offer_price else original_price
                if not base_price:
                    continue
                    
                regular_price = Decimal(str(float(base_price) * 1.1))
                sale_price = None
                
                if offer_price and original_price:
                    sale_price = regular_price
                    regular_price = Decimal(str(float(original_price) * 1.1))
                
                # Get category
                category = category_map.get(category_name)
                if not category:
                    continue
                
                # Generate slug and SKU
                slug = product_name.lower().replace(' ', '-').replace('/', '-').replace('&', 'and').replace('(', '').replace(')', '').replace(',', '').replace('"', '').replace("'", '')
                sku = f'VMS-{index + 1:03d}'
                
                # Determine product type
                product_type = 'single'
                if 'GIFT BOXES' in category_name or 'FAMILY PACK' in category_name:
                    product_type = 'gift_box'
                elif 'PACK' in product_name or 'Pcs' in product_name:
                    product_type = 'box'
                elif 'SET' in product_name or 'COMBO' in product_name:
                    product_type = 'combo'
                
                # Determine safety level
                safety_level = 'medium'
                if any(word in category_name for word in ['BOMB', 'CRACKERS', 'ROCKET']):
                    safety_level = 'high'
                elif any(word in category_name for word in ['SPARKLERS', 'CANDLES', 'TOYS']):
                    safety_level = 'low'
                
                # Create product
                Product.objects.create(
                    name=product_name,
                    slug=slug,
                    sku=sku,
                    category=category,
                    brand=brand,
                    product_type=product_type,
                    safety_level=safety_level,
                    short_description=f'{product_name} - {category_name}',
                    description=f'{product_name} from {category_name}. High quality crackers for celebrations and festivals.',
                    regular_price=regular_price,
                    sale_price=sale_price,
                    stock=random.randint(10, 100),
                    pieces=1,
                    is_active=True,
                    is_featured=random.random() > 0.8,
                    is_new=random.random() > 0.7,
                    is_bestseller=random.random() > 0.85,
                    is_trending=random.random() > 0.6,
                    main_image='images/crackers/placeholder.jpg',
                    image_url='images/crackers/placeholder.jpg',
                    additional_images=[],
                )
                created_count += 1
                
                if created_count % 20 == 0:
                    self.stdout.write(f'Created {created_count} products...')
            
            self.stdout.write(self.style.SUCCESS(f'Successfully restored {created_count} products from Excel'))
            self.stdout.write(f'Total categories: {Category.objects.count()}')
            self.stdout.write(f'Total products: {Product.objects.count()}')
            
        except Exception as e:
            self.stdout.write(self.style.ERROR(f'Error: {e}'))
            import traceback
            traceback.print_exc()